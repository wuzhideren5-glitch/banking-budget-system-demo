from __future__ import annotations

import ast
import asyncio
from contextlib import asynccontextmanager, suppress
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import re
import sqlite3
from typing import Any

import aiosqlite
from fastapi import FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.audit import write_operation_log
from app.agent_graph import AgentGraphService
from app.agent_debug_trace import AgentDebugTraceStore
from app.agent_memory import ConversationMemoryStore
from app.agent_query import ReadOnlySqlExecutor
from app.budget_window import budget_actual_allowed_for_month
from app.config import settings
from app.deepseek_client import DeepseekClient
from app.db_paths import budget_db_path, common_db_path, compare_db_path, list_budget_database_files
from app.formula_refs import extract_formula_codes, formulas_reference_code
from app.feishu_bot import start_feishu_background
from app.init_db import BUDGET_SCHEMA, ensure_databases
from app.knowledge_base import KnowledgeBaseService
from app.months import parse_month_index
from app.passwords import hash_daily_password, verify_daily_password
from app.product_scope_migration import (
    migrate_all_to_single_budget_data,
    migrate_single_to_all_budget_data,
    preview_delete_all_to_single_rows,
    preview_insert_single_to_all_rows,
)
from app.services.compare_export_service import CompareExportService
from app.services.budget_summary_export_service import BudgetSummaryExportService
from app.services.smart_report_service import SmartReportService
from app.services.export_common import (
    autosize_worksheet_columns,
    budget_summary_field_meta,
    build_export_versions_info_text,
    build_export_year_datetime_text,
    normalize_summary_value,
    write_template_pivot_data_area,
)
from app.routers.agent_kb import build_agent_kb_router
from app.routers.agent_runtime import build_agent_runtime_router
from app.routers.agent_debug import build_agent_debug_router
from app.routers.auth import build_auth_router
from app.routers.budget_subject_catalog import build_budget_subject_catalog_router
from app.routers.data_accounts import build_data_accounts_router
from app.routers.dept_catalog import build_dept_catalog_router
from app.routers.budget_assumptions import build_budget_assumptions_router
from app.routers.forecast_workbench import build_forecast_workbench_router
from app.routers.budget_input_import import build_budget_input_import_router
from app.routers.budget_driver import build_budget_driver_router
from app.routers.budget_input_runtime import build_budget_input_runtime_router
from app.routers.fee_actual_import import build_fee_actual_import_router
from app.routers.expense_actual_import import build_expense_actual_import_router
from app.routers.expense_budget_execution import build_expense_budget_execution_router
from app.routers.expense_forecast import build_expense_forecast_router
from app.routers.budget_summary_export import build_budget_summary_export_router
from app.routers.budget_summary_compare import build_budget_summary_compare_router
from app.routers.compare_summary_export import build_compare_summary_export_router
from app.routers.chart_readonly import build_chart_readonly_router
from app.routers.chart_write import build_chart_write_router
from app.routers.health import router as health_router
from app.routers.report_catalog import build_report_catalog_router
from app.routers.smart_reports import build_smart_reports_router
from app.routers.system_admin import build_system_admin_router
from app.routers.system_catalog import build_system_catalog_router
from app.routers.system_edit_show import build_system_edit_show_router
from app.routers.templates import router as templates_router
from app.routers.version_snapshot import build_version_snapshot_router
from app.schemas import (
    DataAccountRow,
    BudgetSummaryExportPivotRequest,
    CompareSummarySyncResult,
    GlobalRefreshAnnualStatus,
    GlobalRefreshStatusResponse,
    ChartVersionItemDto,
    SystemDatabaseRow,
)


@asynccontextmanager
async def lifespan(app: FastAPI):
    global _periodic_global_refresh_task, _periodic_global_refresh_next_run_at
    ensure_databases()
    start_feishu_background(agent_service)
    _periodic_global_refresh_next_run_at = _iso_after_seconds(GLOBAL_REFRESH_INTERVAL_SECONDS)
    _periodic_global_refresh_task = asyncio.create_task(_periodic_global_refresh_loop())
    try:
        yield
    finally:
        if _periodic_global_refresh_task is not None:
            _periodic_global_refresh_task.cancel()
            with suppress(asyncio.CancelledError):
                await _periodic_global_refresh_task
            _periodic_global_refresh_task = None


app = FastAPI(title="Banking Budget API", lifespan=lifespan)
_cors_regex = (settings.cors_origin_regex or "").strip()
app.add_middleware(
    CORSMiddleware,
    allow_origins=[o.strip() for o in settings.cors_origins.split(",") if o.strip()],
    allow_origin_regex=_cors_regex or None,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.include_router(health_router)
app.include_router(templates_router)

_compare_formula_export_jobs: dict[str, dict[str, Any]] = {}
_compare_formula_export_jobs_lock = asyncio.Lock()
repo_root = Path(__file__).resolve().parents[2]
kb_service = KnowledgeBaseService(repo_root)
smart_report_service = SmartReportService(data_dir=settings.data_dir)
query_executor = ReadOnlySqlExecutor(
    budget_db_path=budget_db_path(settings.budget_year),
    common_db_path=common_db_path(),
)
memory_store = ConversationMemoryStore(repo_root)
agent_debug_trace_store = AgentDebugTraceStore(repo_root / "knowledge_base" / "generated" / "agent_llm_events.jsonl")
deepseek_client = DeepseekClient(
    api_key=settings.deepseek_api_key,
    base_url=settings.deepseek_base_url,
    model=settings.deepseek_model,
)
agent_service = AgentGraphService(
    kb_service,
    query_executor=query_executor,
    memory_store=memory_store,
    deepseek_client=deepseek_client,
    debug_trace_store=agent_debug_trace_store,
)
app.include_router(build_agent_kb_router(kb_service))
app.include_router(build_agent_runtime_router(agent_service, memory_store))
app.include_router(build_agent_debug_router(agent_debug_trace_store))
app.include_router(build_budget_assumptions_router())
app.include_router(build_forecast_workbench_router())

SESSION_COOKIE_NAME = "budget_session"
SESSION_TTL_SECONDS = 8 * 60 * 60
GLOBAL_REFRESH_INTERVAL_SECONDS = 10 * 60
BUDGET_GLOBAL_REFRESH_KEY = "global_refresh_time_a"
COMPARE_GLOBAL_REFRESH_KEY = "global_refresh_time_b"
_periodic_global_refresh_task: asyncio.Task | None = None
_periodic_global_refresh_next_run_at: str | None = None
_compare_sync_lock = asyncio.Lock()


def _role_name_from_permission(permission_type: int) -> str:
    if permission_type == 1:
        return "全权管理员"
    if permission_type == 2:
        return "数据录入用户"
    return "数据浏览用户"


def _permission_set(permission_type: int) -> set[int]:
    if permission_type == 1:
        return {1, 2, 3}
    if permission_type == 2:
        return {1, 2}
    return {1}


def _path_required_permission(path: str, method: str) -> int | None:
    if path.startswith("/api/system"):
        return 3
    if path.startswith("/api/data-accounts") or path.startswith("/api/report-accounts"):
        return 3
    if path.startswith("/api/budget-subject-catalog"):
        return 3
    if path.startswith("/api/dept-accounts") or path.startswith("/api/product-types"):
        return 3
    if path.startswith("/api/report-data-mappings") or path.startswith("/api/dept-product-mappings"):
        return 3
    if path.startswith("/api/budget-input"):
        return 2
    if path.startswith("/api/budget-assumptions") or path.startswith("/api/forecast-workbench"):
        return 2
    if path.startswith("/api/driver"):
        return 2
    if path.startswith("/api/fee-actual"):
        return 2
    if path.startswith("/api/expense-actual-import"):
        return 2
    if path.startswith("/api/expense-forecast"):
        return 2
    if path.startswith("/api/budget-summary") or path.startswith("/api/compare-summary"):
        return 1
    if path.startswith("/api/expense-budget-execution/admin"):
        return 3
    if path.startswith("/api/expense-budget-execution"):
        return 1
    if path == "/api/global-refresh-status":
        return 1
    if path.startswith("/api/chart"):
        return 1
    if path.startswith("/api/agent"):
        return 1
    if path.startswith("/api/smart-reports/templates") and method.upper() in {"POST", "PUT", "PATCH", "DELETE"}:
        return 3
    if path.startswith("/api/smart-reports/calc-metrics") and method.upper() in {"POST", "PUT", "PATCH", "DELETE"}:
        return 3
    if path.startswith("/api/smart-reports"):
        return 1
    return None


def _parse_iso_utc(ts: str) -> datetime:
    return datetime.strptime(ts, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)


def _parse_any_ts_utc(ts: str | None) -> datetime | None:
    t = str(ts or "").strip()
    if not t:
        return None
    fmts = (
        "%Y-%m-%dT%H:%M:%SZ",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%dT%H:%M:%S.%fZ",
    )
    for fmt in fmts:
        try:
            return datetime.strptime(t, fmt).replace(tzinfo=timezone.utc)
        except Exception:
            continue
    return None


def _iso_after_seconds(seconds: int) -> str:
    dt = datetime.now(timezone.utc).timestamp() + int(seconds)
    return datetime.fromtimestamp(dt, tz=timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _verify_daily_password(stored_hash: str | None, raw_password: str) -> bool:
    return verify_daily_password(stored_hash, raw_password)


async def _load_session_context(session_id: str | None) -> dict[str, Any] | None:
    if not session_id:
        return None
    now = _iso_now()
    async with aiosqlite.connect(common_db_path()) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            """
            SELECT s.session_id, s.user_id, s.must_change_password, s.expire_time,
                   u.user_name, u.permission_type, u.first_login_flag
            FROM user_sessions s
            JOIN users u ON u.id = s.user_id
            WHERE s.session_id = ?
            """,
            (session_id,),
        )
        row = await cur.fetchone()
        if not row:
            return None
        expire_time = str(row[3])
        if _parse_iso_utc(expire_time) < _parse_iso_utc(now):
            await db.execute("DELETE FROM user_sessions WHERE session_id = ?", (session_id,))
            await db.commit()
            return None
        await db.execute(
            "UPDATE user_sessions SET last_seen_time = ? WHERE session_id = ?",
            (now, session_id),
        )
        await db.commit()
        return {
            "session_id": str(row[0]),
            "user_id": int(row[1]),
            "must_change_password": int(row[2]),
            "expire_time": expire_time,
            "user_name": str(row[4]),
            "permission_type": int(row[5]),
            "first_login_flag": int(row[6]),
        }


@app.middleware("http")
async def auth_middleware(request: Request, call_next):
    path = request.url.path
    if request.method.upper() == "OPTIONS":
        return await call_next(request)
    if not path.startswith("/api"):
        return await call_next(request)

    anonymous_allowed = {
        "/api/health",
        "/api/login",
    }
    if path in anonymous_allowed:
        return await call_next(request)

    session_id = request.cookies.get(SESSION_COOKIE_NAME)
    session_ctx = await _load_session_context(session_id)
    if session_ctx is None:
        if path == "/api/session":
            return JSONResponse({"detail": "未登录"}, status_code=401)
        resp = JSONResponse({"detail": "未登录，请先登录"}, status_code=401)
        resp.delete_cookie(SESSION_COOKIE_NAME)
        return resp

    request.state.current_user = session_ctx
    if session_ctx["must_change_password"] == 1:
        if path not in {"/api/session", "/api/change-password-first-login", "/api/logout"}:
            return JSONResponse({"detail": "首次登录请先修改密码"}, status_code=403)
        return await call_next(request)

    required_permission = _path_required_permission(path, request.method)
    if required_permission is not None:
        allowed = _permission_set(int(session_ctx["permission_type"]))
        if required_permission not in allowed:
            return JSONResponse({"detail": "权限不足"}, status_code=403)
    return await call_next(request)

@app.get("/api/global-refresh-status", response_model=GlobalRefreshStatusResponse)
async def global_refresh_status():
    return await _collect_global_refresh_status()


async def _latest_version() -> tuple[int, str, str]:
    path = budget_db_path(settings.budget_year)
    async with aiosqlite.connect(path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            """
            SELECT version_id, version_name, version_date_time
            FROM version ORDER BY version_id DESC LIMIT 1
            """
        )
        row = await cur.fetchone()
        if not row:
            raise HTTPException(status_code=500, detail="年度库缺少 version 记录")
        return int(row[0]), str(row[1]), str(row[2])


async def _editable_context() -> tuple[Path, int, int]:
    """Current editable (db path, year, version_id) from edit_show_version sign=0."""
    async with aiosqlite.connect(common_db_path()) as cdb:
        await cdb.execute("PRAGMA foreign_keys = ON")
        cur = await cdb.execute(
            """
            SELECT d.data_file_name, d.year, e.version_id
            FROM edit_show_version e
            JOIN databases d ON d.id = e.data_file_id
            WHERE e.edit_show_sign = 0
            LIMIT 1
            """
        )
        row = await cur.fetchone()
    if row:
        data_file_name = str(row[0])
        year = int(row[1])
        version_id = int(row[2])
        return settings.data_dir / data_file_name, year, version_id
    # Fallback: legacy default behavior.
    path = budget_db_path(settings.budget_year)
    vid, _vn, _vdt = await _latest_version()
    return path, settings.budget_year, vid


async def _latest_version_in_path(path: Path) -> tuple[int, str, str]:
    async with aiosqlite.connect(path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            """
            SELECT version_id, version_name, version_date_time
            FROM version ORDER BY version_id DESC LIMIT 1
            """
        )
        row = await cur.fetchone()
    if not row:
        raise HTTPException(status_code=500, detail="年度库缺少 version 记录")
    return int(row[0]), str(row[1]), str(row[2])


async def _try_latest_version_id() -> int | None:
    """Latest version_id in the current budget year DB, or None if file/table/row missing.

    Does not create a new DB file (avoids empty budget_*.db on first connect).
    Used after product-scope migration so we skip formula recalc instead of HTTP 500
    when the current-year file is absent or has no version row yet.
    """
    path = budget_db_path(settings.budget_year)
    if not path.exists():
        return None
    async with aiosqlite.connect(path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='version'"
        )
        if not await cur.fetchone():
            return None
        cur = await db.execute(
            "SELECT version_id FROM version ORDER BY version_id DESC LIMIT 1"
        )
        row = await cur.fetchone()
        if not row:
            return None
        return int(row[0])


async def _last_calc_time(path: Path | None = None) -> str | None:
    compare_path = compare_db_path()
    if compare_path.exists():
        async with aiosqlite.connect(compare_path) as db:
            await db.execute("PRAGMA foreign_keys = ON")
            await db.execute(
                """
                CREATE TABLE IF NOT EXISTS settings (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  setting_key TEXT NOT NULL UNIQUE,
                  setting_value TEXT NOT NULL
                )
                """
            )
            cur = await db.execute(
                "SELECT setting_value FROM settings WHERE setting_key = ? LIMIT 1",
                (COMPARE_GLOBAL_REFRESH_KEY,),
            )
            row = await cur.fetchone()
            if row and row[0]:
                return str(row[0])
    budget_path = path if path is not None else budget_db_path(settings.budget_year)
    async with aiosqlite.connect(budget_path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            "SELECT MAX(update_time) FROM budget_summary WHERE update_time IS NOT NULL"
        )
        row = await cur.fetchone()
        if row and row[0]:
            return str(row[0])
    return None


async def _active_session_count() -> int:
    now = _iso_now()
    async with aiosqlite.connect(common_db_path()) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        await db.execute(
            "DELETE FROM user_sessions WHERE expire_time < ?",
            (now,),
        )
        cur = await db.execute(
            "SELECT COUNT(*) FROM user_sessions WHERE expire_time >= ?",
            (now,),
        )
        row = await cur.fetchone()
        await db.commit()
    return int(row[0] or 0) if row else 0


async def _ensure_budget_settings_table(budget_path: Path) -> None:
    async with aiosqlite.connect(budget_path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        await db.execute(
            """
            CREATE TABLE IF NOT EXISTS settings (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              setting_key TEXT NOT NULL UNIQUE,
              setting_value TEXT NOT NULL
            )
            """
        )
        await db.commit()


async def _get_budget_refresh_time_a(budget_path: Path) -> str | None:
    await _ensure_budget_settings_table(budget_path)
    async with aiosqlite.connect(budget_path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            "SELECT setting_value FROM settings WHERE setting_key = ? LIMIT 1",
            (BUDGET_GLOBAL_REFRESH_KEY,),
        )
        row = await cur.fetchone()
    return str(row[0]) if row and row[0] is not None else None


async def _set_budget_refresh_time_a(budget_path: Path, ts: str) -> None:
    await _ensure_budget_settings_table(budget_path)
    async with aiosqlite.connect(budget_path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        await db.execute(
            """
            INSERT INTO settings(setting_key, setting_value)
            VALUES (?, ?)
            ON CONFLICT(setting_key) DO UPDATE SET setting_value = excluded.setting_value
            """,
            (BUDGET_GLOBAL_REFRESH_KEY, ts),
        )
        await db.commit()


async def _ensure_compare_settings_table() -> None:
    async with aiosqlite.connect(compare_db_path()) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        await db.execute(
            """
            CREATE TABLE IF NOT EXISTS settings (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              setting_key TEXT NOT NULL UNIQUE,
              setting_value TEXT NOT NULL
            )
            """
        )
        await db.commit()


async def _get_compare_refresh_time_b() -> str | None:
    await _ensure_compare_settings_table()
    async with aiosqlite.connect(compare_db_path()) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            "SELECT setting_value FROM settings WHERE setting_key = ? LIMIT 1",
            (COMPARE_GLOBAL_REFRESH_KEY,),
        )
        row = await cur.fetchone()
    return str(row[0]) if row and row[0] is not None else None


async def _set_compare_refresh_time_b(ts: str) -> None:
    await _ensure_compare_settings_table()
    async with aiosqlite.connect(compare_db_path()) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        await db.execute(
            """
            INSERT INTO settings(setting_key, setting_value)
            VALUES (?, ?)
            ON CONFLICT(setting_key) DO UPDATE SET setting_value = excluded.setting_value
            """,
            (COMPARE_GLOBAL_REFRESH_KEY, ts),
        )
        await db.commit()


async def _max_budget_data_update_time(budget_path: Path) -> str | None:
    async with aiosqlite.connect(budget_path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            "SELECT MAX(update_time) FROM budget_data WHERE update_time IS NOT NULL"
        )
        row = await cur.fetchone()
    return str(row[0]) if row and row[0] is not None else None


async def _rebuild_budget_summary_all_versions(budget_path: Path) -> int:
    async with aiosqlite.connect(budget_path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            "SELECT version_id FROM version ORDER BY version_id"
        )
        version_ids = [int(r[0]) for r in await cur.fetchall()]
    inserted = 0
    for vid in version_ids:
        inserted += int(await _rebuild_budget_summary_for_version(vid, budget_path))
    return inserted


async def _recalculate_dirty_budget_data_formulas_for_year(budget_path: Path, budget_year: int) -> int:
    async with aiosqlite.connect(common_db_path()) as cdb:
        await cdb.execute("PRAGMA foreign_keys = ON")
        cur = await cdb.execute(
            """
            SELECT data_acct_code, product_code, product_codes, budget_formula, actual_formula
            FROM data_account
            """
        )
        rows = await cur.fetchall()
    formula_meta: dict[str, dict[str, Any]] = {}
    for r in rows:
        code = str(r[0])
        formula_meta[code] = {
            "product_code": str(r[1] or "").strip().upper(),
            "budget_formula": _normalize_formula(r[3]),
            "actual_formula": _normalize_formula(r[4]),
        }
    async with aiosqlite.connect(budget_path) as bdb:
        await bdb.execute("PRAGMA foreign_keys = ON")
        cur = await bdb.execute(
            """
            SELECT DISTINCT version_id, budget_actual, data_acct_code
            FROM budget_data
            WHERE need_calc = 1
            """
        )
        dirty_rows = await cur.fetchall()
    recalculated = 0
    for version_id, budget_actual, data_code_raw in dirty_rows:
        data_code = str(data_code_raw or "").strip().upper()
        if not data_code:
            continue
        meta = formula_meta.get(data_code)
        if not meta:
            continue
        ba = int(budget_actual or 0)
        formula = str(meta["budget_formula"] if ba == 0 else meta["actual_formula"] or "").strip()
        if not formula:
            continue
        _pcs = meta.get("product_codes")
        # 三态：'all'=全部产品, ''=公司级, 'Z01,Z02'=指定产品
        if _pcs is None or (_pcs is not None and str(_pcs).upper().strip() == 'ALL'):
            recalculated += int(
                await _recalculate_data_account_formula_all_products(
                    data_acct_code=data_code,
                    formula=formula,
                    version_id=int(version_id),
                    budget_actual=ba,
                    budget_path=budget_path,
                    budget_year=budget_year,
                )
            )
        elif str(_pcs).strip() == "":
            # 公司级科目，跳过或按需处理
            continue
        else:
            product_code = str(meta.get("product_code") or "").strip().upper()
            if not product_code:
                continue
            recalculated += int(
                await _recalculate_data_account_formula(
                    data_acct_code=data_code,
                    formula=formula,
                    version_id=int(version_id),
                    budget_actual=ba,
                    product_code=product_code,
                    budget_path=budget_path,
                    budget_year=budget_year,
                )
            )
    return recalculated


async def _run_periodic_global_refresh_once() -> None:
    active_sessions = await _active_session_count()
    if active_sessions <= 0:
        return
    annual_a_times: list[datetime] = []
    for budget_path in list_budget_database_files():
        parsed_year = _parse_year_from_budget_filename(budget_path.name)
        if parsed_year is None:
            continue
        year = int(parsed_year)
        await _recalculate_dirty_budget_data_formulas_for_year(budget_path, year)
        last_a_raw = await _get_budget_refresh_time_a(budget_path)
        last_a = _parse_any_ts_utc(last_a_raw)
        latest_budget_data_raw = await _max_budget_data_update_time(budget_path)
        latest_budget_data = _parse_any_ts_utc(latest_budget_data_raw)
        if latest_budget_data is None:
            if last_a is not None:
                annual_a_times.append(last_a)
            continue
        if last_a is None or latest_budget_data > last_a:
            await _rebuild_budget_summary_all_versions(budget_path)
            refreshed = _iso_now()
            await _set_budget_refresh_time_a(budget_path, refreshed)
            annual_a_times.append(_parse_iso_utc(refreshed))
        else:
            annual_a_times.append(last_a)
    if not annual_a_times:
        return
    max_a = max(annual_a_times)
    b_raw = await _get_compare_refresh_time_b()
    b_time = _parse_any_ts_utc(b_raw)
    if b_time is None or b_time < max_a:
        await _sync_compare_budget_summary(trigger_source="auto_scheduler_10min")


async def _periodic_global_refresh_loop() -> None:
    global _periodic_global_refresh_next_run_at
    while True:
        try:
            await _run_periodic_global_refresh_once()
        except Exception:
            # 后台任务异常不应中断 API 服务；错误细节由 compare/job log 与调用方排障。
            pass
        _periodic_global_refresh_next_run_at = _iso_after_seconds(GLOBAL_REFRESH_INTERVAL_SECONDS)
        await asyncio.sleep(GLOBAL_REFRESH_INTERVAL_SECONDS)


async def _collect_global_refresh_status() -> GlobalRefreshStatusResponse:
    annual_items: list[GlobalRefreshAnnualStatus] = []
    for bpath in list_budget_database_files():
        parsed_year = _parse_year_from_budget_filename(bpath.name)
        if parsed_year is None:
            continue
        year = int(parsed_year)
        a = await _get_budget_refresh_time_a(bpath)
        annual_items.append(
            GlobalRefreshAnnualStatus(
                data_file_name=bpath.name,
                year=int(year),
                refresh_time_a=a,
            )
        )
    annual_items.sort(key=lambda x: (-x.year, x.data_file_name))
    b = await _get_compare_refresh_time_b()
    return GlobalRefreshStatusResponse(
        annual_items=annual_items,
        compare_refresh_time_b=b,
        next_planned_refresh_time_c=_periodic_global_refresh_next_run_at,
    )


async def _is_formula_locked_data_account(data_acct_code: str, budget_actual: int) -> bool:
    code = data_acct_code.strip().upper()
    path = common_db_path()
    async with aiosqlite.connect(path) as db:
        cur = await db.execute(
            """
            SELECT budget_formula, actual_formula
            FROM data_account
            WHERE data_acct_code = ?
            """,
            (code,),
        )
        row = await cur.fetchone()
    if not row:
        return False
    formula = _normalize_formula(row[0] if budget_actual == 0 else row[1])
    return bool(formula)


def _month_index(month_label: str) -> int:
    return parse_month_index(month_label)


def _is_month_editable(current_month: int, budget_actual: int, month_index: int) -> bool:
    if budget_actual == 0:
        return month_index >= current_month
    return month_index < current_month


def _budget_actual_allowed_for_month(budget_actual: int, month: int, current_month: int) -> bool:
    """当前月份窗口 X 下，日历月 month 是否允许存在该 budget_actual（0=预算 1=实际）。"""
    return budget_actual_allowed_for_month(budget_actual, month, current_month)


async def _purge_disallowed_budget_data_for_version(
    bdb: aiosqlite.Connection,
    version_id: int,
    current_month: int,
    period_month_map: dict[int, int],
) -> None:
    """删除违反「X 前仅实际、X 及后仅预算」以及 X=1/X=13 特例的 budget_data 行。"""
    if not period_month_map:
        return
    x = max(1, min(13, current_month))
    if x == 13:
        pids = [pid for pid, m in period_month_map.items() if 1 <= m <= 12]
        if pids:
            ph = ",".join(["?"] * len(pids))
            await bdb.execute(
                f"DELETE FROM budget_data WHERE version_id = ? AND budget_actual = 0 AND period_id IN ({ph})",
                (version_id, *pids),
            )
        return
    if x == 1:
        pids = [pid for pid, m in period_month_map.items() if 1 <= m <= 12]
        if pids:
            ph = ",".join(["?"] * len(pids))
            await bdb.execute(
                f"DELETE FROM budget_data WHERE version_id = ? AND budget_actual = 1 AND period_id IN ({ph})",
                (version_id, *pids),
            )
        return
    pids_budget_bad = [pid for pid, m in period_month_map.items() if 1 <= m < x]
    pids_actual_bad = [pid for pid, m in period_month_map.items() if x <= m <= 12]
    if pids_budget_bad:
        ph = ",".join(["?"] * len(pids_budget_bad))
        await bdb.execute(
            f"DELETE FROM budget_data WHERE version_id = ? AND budget_actual = 0 AND period_id IN ({ph})",
            (version_id, *pids_budget_bad),
        )
    if pids_actual_bad:
        ph = ",".join(["?"] * len(pids_actual_bad))
        await bdb.execute(
            f"DELETE FROM budget_data WHERE version_id = ? AND budget_actual = 1 AND period_id IN ({ph})",
            (version_id, *pids_actual_bad),
        )


async def _get_version_current_month(version_id: int, budget_path: Path | None = None) -> int:
    path = budget_path if budget_path is not None else budget_db_path(settings.budget_year)
    async with aiosqlite.connect(path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            "SELECT current_month FROM version WHERE version_id = ?",
            (version_id,),
        )
        row = await cur.fetchone()
    if not row:
        raise HTTPException(status_code=400, detail=f"版本 {version_id} 不存在")
    current_month = int(row[0] or 1)
    if current_month < 1 or current_month > 13:
        return 1
    return current_month


async def _get_period_month_index(period_id: int) -> int:
    path = common_db_path()
    async with aiosqlite.connect(path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            "SELECT month FROM period WHERE period_id = ?",
            (period_id,),
        )
        row = await cur.fetchone()
    if not row:
        raise HTTPException(status_code=400, detail=f"期间 {period_id} 不存在")
    month_idx = _month_index(str(row[0] or ""))
    if month_idx < 1 or month_idx > 12:
        raise HTTPException(status_code=400, detail=f"期间 {period_id} 月份无效")
    return month_idx


async def _ensure_budget_input_period_editable(
    version_id: int,
    budget_actual: int,
    period_id: int,
    budget_path: Path | None = None,
) -> None:
    current_month = await _get_version_current_month(version_id, budget_path)
    month_idx = await _get_period_month_index(period_id)
    if not _is_month_editable(current_month, budget_actual, month_idx):
        data_kind = "预算值" if budget_actual == 0 else "实际值"
        raise HTTPException(
            status_code=400,
            detail=f"当前版本月份窗口限制：{data_kind}不允许写入 {month_idx} 月（current_month={current_month}）",
        )


def _normalize_formula(value: str | None) -> str:
    return (value or "").strip()


async def _load_data_account_scope_map(db: aiosqlite.Connection) -> dict[str, bool]:
    cur = await db.execute("SELECT data_acct_code, product_codes FROM data_account")
    rows = await cur.fetchall()
    # 三态：'all'=所有产品(True), ''/有值=指定产品或公司级(False)
    return {str(r[0]): r[1] is None or (r[1] is not None and str(r[1]).upper() == 'ALL') for r in rows}


def _validate_formula_reference_scope(
    *,
    formula: str | None,
    target_is_all: bool,
    scope_by_code: dict[str, bool],
    formula_label: str,
) -> None:
    normalized = _normalize_formula(formula)
    if not normalized:
        return
    ref_codes = sorted(extract_formula_codes(normalized))
    if not ref_codes:
        return
    missing_codes = [code for code in ref_codes if code not in scope_by_code]
    if missing_codes:
        raise HTTPException(
            status_code=400,
            detail=f"{formula_label}引用了不存在的数据科目：{', '.join(missing_codes)}",
        )
    if target_is_all:
        disallowed_codes = [code for code in ref_codes if not scope_by_code.get(code, False)]
        if disallowed_codes:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"{formula_label}仅可引用“适用所有产品科目”。"
                    f"以下被引用科目不满足约束：{', '.join(disallowed_codes)}"
                ),
            )


def _prepare_formula_expression(formula: str | None) -> str:
    expr = _normalize_formula(formula)
    if not expr:
        return ""
    # 支持历史公式写法：<A1200 科目名称> -> A1200
    expr = re.sub(r"<\s*([A-Z]\d{4})[^>]*>", r"\1", expr)
    # 兼容全角标点与常见中文运算符。
    translate_map = str.maketrans({
        "（": "(",
        "）": ")",
        "，": ",",
        "＋": "+",
        "－": "-",
        "×": "*",
        "÷": "/",
    })
    return expr.translate(translate_map)


def _eval_formula_ast(node: ast.AST, values: dict[str, float]) -> float:
    if isinstance(node, ast.Expression):
        return _eval_formula_ast(node.body, values)
    if isinstance(node, ast.Constant):
        if isinstance(node.value, (int, float)):
            return float(node.value)
        raise ValueError("公式包含非法常量")
    if isinstance(node, ast.Name):
        return float(values.get(node.id, 0.0))
    if isinstance(node, ast.BinOp):
        left = _eval_formula_ast(node.left, values)
        right = _eval_formula_ast(node.right, values)
        if isinstance(node.op, ast.Add):
            return left + right
        if isinstance(node.op, ast.Sub):
            return left - right
        if isinstance(node.op, ast.Mult):
            return left * right
        if isinstance(node.op, ast.Div):
            if right == 0:
                raise ZeroDivisionError("division by zero")
            return left / right
        raise ValueError("公式仅支持 + - * / 运算")
    if isinstance(node, ast.UnaryOp):
        val = _eval_formula_ast(node.operand, values)
        if isinstance(node.op, ast.UAdd):
            return +val
        if isinstance(node.op, ast.USub):
            return -val
        raise ValueError("公式一元运算不合法")
    if isinstance(node, ast.Call):
        if node.keywords:
            raise ValueError("函数调用不支持命名参数")
        if not isinstance(node.func, ast.Name):
            raise ValueError("函数调用不合法")
        fn = node.func.id.upper()
        args = [_eval_formula_ast(arg, values) for arg in node.args]
        if fn == "SUM":
            return float(sum(args))
        if fn == "AVG":
            return float(sum(args) / len(args)) if args else 0.0
        if fn == "MAX":
            return float(max(args)) if args else 0.0
        if fn == "MIN":
            return float(min(args)) if args else 0.0
        raise ValueError("仅支持 SUM/AVG/MAX/MIN 函数")
    raise ValueError("公式语法不合法")


def _calculate_formula_value(formula: str | None, values: dict[str, float]) -> float:
    value, _ = _try_calculate_formula_value(formula, values)
    return value


def _try_calculate_formula_value(
    formula: str | None, values: dict[str, float]
) -> tuple[float, str | None]:
    expression = _prepare_formula_expression(formula)
    if not expression:
        return 0.0, None
    try:
        parsed = ast.parse(expression, mode="eval")
        return float(_eval_formula_ast(parsed, values)), None
    except ZeroDivisionError:
        return 0.0, "#DIV/0!"
    except Exception:
        return 0.0, "#ERROR!"


def _normalize_formula_ref_value(raw_value: float, value_type: str | None) -> float:
    if (value_type or "") == "百分比" and abs(raw_value) > 1:
        # Backward compatibility: historical percentage may be stored as 5 for 5%.
        return raw_value / 100.0
    return raw_value


async def _period_ids_for_year(budget_year: int) -> list[int]:
    year_label = f"Y{budget_year}"
    path = common_db_path()
    async with aiosqlite.connect(path) as db:
        cur = await db.execute(
            """
            SELECT period_id
            FROM period
            WHERE year = ?
            ORDER BY period_id
            """,
            (year_label,),
        )
        rows = await cur.fetchall()
    return [int(r[0]) for r in rows]


async def _recalculate_data_account_formula(
    data_acct_code: str,
    formula: str | None,
    version_id: int,
    budget_actual: int,
    product_code: str,
    *,
    budget_path: Path | None = None,
    budget_year: int | None = None,
) -> int:
    resolved_budget_year = budget_year if budget_year is not None else settings.budget_year
    period_ids = await _period_ids_for_year(resolved_budget_year)
    if not period_ids:
        return 0

    ref_codes = sorted(extract_formula_codes(formula))
    resolved_budget_path = budget_path if budget_path is not None else budget_db_path(resolved_budget_year)
    common_path = common_db_path()
    pc_norm = product_code.strip().upper()
    value_map: dict[tuple[str, int], float] = {}
    ref_value_type_map: dict[str, str] = {}
    if ref_codes:
        placeholders = ",".join(["?"] * len(ref_codes))
        async with aiosqlite.connect(common_path) as cdb:
            await cdb.execute("PRAGMA foreign_keys = ON")
            cur_types = await cdb.execute(
                f"""
                SELECT data_acct_code, value_type
                FROM data_account
                WHERE data_acct_code IN ({placeholders})
                """,
                tuple(ref_codes),
            )
            ref_value_type_map = {str(r[0]): str(r[1] or "") for r in await cur_types.fetchall()}
    async with aiosqlite.connect(resolved_budget_path) as db:
        await db.execute("PRAGMA foreign_keys = ON")

        if ref_codes:
            period_placeholders = ",".join(["?"] * len(period_ids))
            code_placeholders = ",".join(["?"] * len(ref_codes))
            cur = await db.execute(
                f"""
                SELECT data_acct_code, period_id, value
                FROM budget_data
                WHERE version_id = ?
                  AND budget_actual = ?
                  AND product_code = ?
                  AND period_id IN ({period_placeholders})
                  AND data_acct_code IN ({code_placeholders})
                """,
                (version_id, budget_actual, pc_norm, *period_ids, *ref_codes),
            )
            for r in await cur.fetchall():
                value_map[(str(r[0]), int(r[1]))] = float(r[2] or 0.0)

        now = _iso_now()
        for period_id in period_ids:
            refs_for_period = {
                code: _normalize_formula_ref_value(
                    value_map.get((code, period_id), 0.0),
                    ref_value_type_map.get(code, ""),
                )
                for code in ref_codes
            }
            value = _calculate_formula_value(formula, refs_for_period)
            await db.execute(
                """
                INSERT INTO budget_data (
                  data_acct_code, product_code, period_id, budget_actual, version_id, value, need_calc, create_time, update_time
                ) VALUES (?, ?, ?, ?, ?, ?, 0, ?, ?)
                ON CONFLICT(data_acct_code, product_code, period_id, version_id, budget_actual)
                DO UPDATE SET value = excluded.value, need_calc = 0, update_time = excluded.update_time
                """,
                (data_acct_code, pc_norm, period_id, budget_actual, version_id, value, now, now),
            )
        await db.commit()
    return len(period_ids)


async def _recalculate_data_account_formula_all_products(
    data_acct_code: str,
    formula: str | None,
    version_id: int,
    budget_actual: int,
    *,
    budget_path: Path | None = None,
    budget_year: int | None = None,
) -> int:
    common_path = common_db_path()
    code_u = data_acct_code.strip().upper()
    async with aiosqlite.connect(common_path) as cdb:
        await cdb.execute("PRAGMA foreign_keys = ON")
        cur = await cdb.execute(
            "SELECT product_codes, product_code FROM data_account WHERE data_acct_code = ?",
            (code_u,),
        )
        row = await cur.fetchone()
    if not row:
        return 0
    # 三态：product_codes 'all'=所有产品, 有值=指定产品, ''=公司级
    applies_all = row[0] is None or (row[0] is not None and str(row[0]).upper() == 'ALL')
    single_pc = str(row[1]).strip().upper() if row[1] else ""
    total = 0
    if applies_all:
        async with aiosqlite.connect(common_path) as cdb:
            await cdb.execute("PRAGMA foreign_keys = ON")
            cur = await cdb.execute("SELECT product_code FROM product_type ORDER BY product_code")
            products = [str(r[0]) for r in await cur.fetchall()]
        if not products:
            return 0
        for p in products:
            total += await _recalculate_data_account_formula(
                code_u,
                formula,
                version_id,
                budget_actual,
                p,
                budget_path=budget_path,
                budget_year=budget_year,
            )
    else:
        if not single_pc:
            return 0
        total += await _recalculate_data_account_formula(
            code_u,
            formula,
            version_id,
            budget_actual,
            single_pc,
            budget_path=budget_path,
            budget_year=budget_year,
        )
    return total


async def _set_data_account_need_calc(code: str, flag: int) -> None:
    path = common_db_path()
    async with aiosqlite.connect(path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        await db.execute(
            "UPDATE data_account SET need_calc = ? WHERE data_acct_code = ?",
            (1 if flag else 0, code),
        )
        await db.commit()


async def _recalculate_product_formula_rows(
    product_code: str,
    version_id: int,
    budget_actual: int,
    *,
    budget_path: Path | None = None,
    budget_year: int | None = None,
) -> int:
    common_path = common_db_path()
    formula_col = "budget_formula" if budget_actual == 0 else "actual_formula"
    async with aiosqlite.connect(common_path) as cdb:
        await cdb.execute("PRAGMA foreign_keys = ON")
        cur = await cdb.execute(
            f"""
            SELECT data_acct_code, {formula_col}
            FROM data_account
            WHERE product_code = ? OR (product_codes IS NULL OR product_codes = '')
            """,
            (product_code.strip().upper(),),
        )
        rows = await cur.fetchall()
    formulas = [(str(r[0]), _normalize_formula(r[1])) for r in rows if _normalize_formula(r[1])]
    if not formulas:
        return 0
    recalculated = 0
    pc = product_code.strip().upper()
    for code, formula in formulas:
        recalculated += await _recalculate_data_account_formula(
            data_acct_code=code,
            formula=formula,
            version_id=version_id,
            budget_actual=budget_actual,
            product_code=pc,
            budget_path=budget_path,
            budget_year=budget_year,
        )
    return recalculated


async def _set_budget_data_need_calc_for_cells(
    cells: list[tuple[str, str, int, int, int]],
    flag: int,
    budget_path: Path | None = None,
) -> None:
    if not cells:
        return
    path = budget_path if budget_path is not None else budget_db_path(settings.budget_year)
    now = _iso_now()
    async with aiosqlite.connect(path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        for data_acct_code, product_code, period_id, version_id, budget_actual in cells:
            await db.execute(
                """
                INSERT INTO budget_data (
                  data_acct_code, product_code, period_id, budget_actual, version_id, value, need_calc, create_time, update_time
                ) VALUES (?, ?, ?, ?, ?, 0, ?, ?, ?)
                ON CONFLICT(data_acct_code, product_code, period_id, version_id, budget_actual)
                DO UPDATE SET need_calc = excluded.need_calc, update_time = excluded.update_time
                """,
                (
                    data_acct_code,
                    product_code.strip().upper(),
                    period_id,
                    budget_actual,
                    version_id,
                    1 if flag else 0,
                    now,
                    now,
                ),
            )
        await db.commit()


def _build_report_path(
    report_code: str, report_name_map: dict[str, str], report_parent_map: dict[str, str | None]
) -> list[str]:
    path: list[str] = []
    seen: set[str] = set()
    cur = report_code
    while cur and cur not in seen:
        seen.add(cur)
        name = report_name_map.get(cur, "")
        path.append(f"{cur} {name}".strip())
        cur = report_parent_map.get(cur) or ""
    path.reverse()
    return path


def _iso_now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _build_fixed_levels(path: list[str], max_levels: int) -> list[str | None]:
    levels: list[str | None] = [None] * max_levels
    for i, token in enumerate(path[:max_levels]):
        levels[i] = token
    return levels


app.include_router(
    build_auth_router(
        settings=settings,
        session_cookie_name=SESSION_COOKIE_NAME,
        session_ttl_seconds=SESSION_TTL_SECONDS,
        verify_daily_password=_verify_daily_password,
        hash_daily_password=hash_daily_password,
        validate_password_policy=lambda password: _validate_password_policy(password),
        role_name_from_permission=_role_name_from_permission,
        iso_now=_iso_now,
        editable_context_provider=_editable_context,
        latest_version_in_path_provider=_latest_version_in_path,
        last_calc_time_provider=_last_calc_time,
    )
)
app.include_router(
    build_version_snapshot_router(
        lambda data_file_name, version_id: _fetch_version_name_and_current_month_from_file(data_file_name, version_id)
    )
)
app.include_router(
    build_smart_reports_router(
        service=smart_report_service,
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
    )
)


async def _fetch_version_name_and_current_month_from_file(data_file_name: str, version_id: int) -> tuple[str, int]:
    budget_path = settings.data_dir / data_file_name
    if not budget_path.exists():
        return (f"V{version_id}", 1)
    async with aiosqlite.connect(budget_path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute("PRAGMA table_info(version)")
        cols = {str(r[1]) for r in await cur.fetchall()}
        has_current_month = "current_month" in cols
        sql = (
            "SELECT version_name, current_month FROM version WHERE version_id = ?"
            if has_current_month
            else "SELECT version_name, 1 AS current_month FROM version WHERE version_id = ?"
        )
        cur = await db.execute(sql, (version_id,))
        row = await cur.fetchone()
    if not row or row[0] is None:
        return (f"V{version_id}", 1)
    return (str(row[0]), int(row[1] or 1))


async def _chart_version_options() -> list[ChartVersionItemDto]:
    async with aiosqlite.connect(common_db_path()) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            """
            SELECT e.edit_show_sign, d.id, d.data_file_name, d.year, e.version_id
            FROM edit_show_version e
            JOIN databases d ON d.id = e.data_file_id
            WHERE e.edit_show_sign BETWEEN 1 AND 5
            ORDER BY e.edit_show_sign
            """
        )
        selected_rows = await cur.fetchall()
        if not selected_rows:
            cur = await db.execute(
                """
                SELECT 1 AS edit_show_sign, id, data_file_name, year, NULL as version_id
                FROM databases
                ORDER BY year DESC, id DESC
                """
            )
            selected_rows = await cur.fetchall()
    options: list[ChartVersionItemDto] = []
    for db_row in selected_rows:
        show_level = int(db_row[0] or 0)
        data_file_id = int(db_row[1])
        data_file_name = str(db_row[2])
        year = int(db_row[3])
        selected_version_id = int(db_row[4]) if db_row[4] is not None else None
        budget_path = settings.data_dir / data_file_name
        if not budget_path.exists():
            continue
        async with aiosqlite.connect(budget_path) as bdb:
            await bdb.execute("PRAGMA foreign_keys = ON")
            if selected_version_id is not None:
                cur_versions = await bdb.execute(
                    """
                    SELECT version_id, version_name, current_month
                    FROM version
                    WHERE version_id = ?
                    ORDER BY version_id
                    """,
                    (selected_version_id,),
                )
            else:
                cur_versions = await bdb.execute(
                    """
                    SELECT version_id, version_name, current_month
                    FROM version
                    ORDER BY version_id
                    """
                )
            version_rows = await cur_versions.fetchall()
        for vr in version_rows:
            version_name = str(vr[1])
            if show_level >= 1:
                version_name = f"L{show_level} {version_name}"
            options.append(
                ChartVersionItemDto(
                    data_file_id=data_file_id,
                    data_file_name=data_file_name,
                    year=year,
                    version_id=int(vr[0]),
                    version_name=version_name,
                    current_month=int(vr[2] or 1),
                )
            )
    return options


app.include_router(build_chart_readonly_router(_chart_version_options))


app.include_router(
    build_chart_write_router(
        settings=settings,
        chart_version_options_provider=_chart_version_options,
        extract_data_acct_code_from_name=lambda data_code_name: _extract_data_acct_code_from_name(data_code_name),
    )
)


app.include_router(
    build_budget_input_import_router(
        editable_context_provider=lambda: _editable_context(),
        recalculate_product_formula_rows=lambda *args, **kwargs: _recalculate_product_formula_rows(*args, **kwargs),
        get_year_period_months=lambda year: _get_year_period_months(year),
        get_version_current_month=lambda version_id, budget_path: _get_version_current_month(version_id, budget_path),
        purge_disallowed_budget_data_for_version=(
            lambda bdb, version_id, current_month, period_month_map: _purge_disallowed_budget_data_for_version(
                bdb, version_id, current_month, period_month_map
            )
        ),
        normalize_cell=lambda value: _normalize_cell(value),
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
    )
)


app.include_router(
    build_fee_actual_import_router(
        editable_context_provider=lambda: _editable_context(),
        recalculate_product_formula_rows=lambda *args, **kwargs: _recalculate_product_formula_rows(*args, **kwargs),
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
    )
)


app.include_router(
    build_budget_driver_router(
        editable_context_provider=lambda: _editable_context(),
        recalculate_product_formula_rows=lambda *args, **kwargs: _recalculate_product_formula_rows(*args, **kwargs),
        rebuild_budget_summary=lambda version_id, budget_path: _rebuild_budget_summary_for_version(version_id, budget_path),
        get_year_period_months=lambda year: _get_year_period_months(year),
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
    )
)


app.include_router(
    build_budget_input_runtime_router(
        editable_context_provider=lambda: _editable_context(),
        get_version_current_month=lambda version_id, budget_path: _get_version_current_month(version_id, budget_path),
        get_year_period_months=lambda year: _get_year_period_months(year),
        purge_disallowed_budget_data_for_version=(
            lambda bdb, version_id, current_month, period_month_map: _purge_disallowed_budget_data_for_version(
                bdb, version_id, current_month, period_month_map
            )
        ),
        normalize_formula=lambda value: _normalize_formula(value),
        month_index=lambda month_label: _month_index(month_label),
        is_month_editable=lambda current_month, budget_actual, month_index: _is_month_editable(current_month, budget_actual, month_index),
        budget_actual_allowed_for_month=(
            lambda budget_actual, month, current_month: _budget_actual_allowed_for_month(budget_actual, month, current_month)
        ),
        normalize_formula_ref_value=(
            lambda raw_value, value_type: _normalize_formula_ref_value(raw_value, value_type)
        ),
        try_calculate_formula_value=(
            lambda formula, values: _try_calculate_formula_value(formula, values)
        ),
        build_report_path=(
            lambda report_code, report_name_map, report_parent_map: _build_report_path(
                report_code, report_name_map, report_parent_map
            )
        ),
        is_formula_locked_data_account=(
            lambda data_acct_code, budget_actual: _is_formula_locked_data_account(data_acct_code, budget_actual)
        ),
        ensure_budget_input_period_editable=(
            lambda version_id, budget_actual, period_id, budget_path: _ensure_budget_input_period_editable(
                version_id=version_id,
                budget_actual=budget_actual,
                period_id=period_id,
                budget_path=budget_path,
            )
        ),
        set_budget_data_need_calc_for_cells=(
            lambda edited_cells, need_calc, budget_path: _set_budget_data_need_calc_for_cells(
                edited_cells, need_calc, budget_path
            )
        ),
        recalculate_product_formula_rows=(
            lambda *args, **kwargs: _recalculate_product_formula_rows(*args, **kwargs)
        ),
        sync_compare_budget_summary=lambda **kwargs: _sync_compare_budget_summary(**kwargs),
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
    )
)


async def _rebuild_budget_summary_for_version(
    version_id: int,
    budget_path: Path | None = None,
) -> int:
    common_path = common_db_path()
    resolved_budget_path = budget_path if budget_path is not None else budget_db_path(settings.budget_year)
    async with aiosqlite.connect(common_path) as cdb, aiosqlite.connect(resolved_budget_path) as bdb:
        await cdb.execute("PRAGMA foreign_keys = ON")
        await bdb.execute("PRAGMA foreign_keys = ON")

        cur_ver = await bdb.execute(
            "SELECT version_name, current_month FROM version WHERE version_id = ?",
            (version_id,),
        )
        vrow = await cur_ver.fetchone()
        if not vrow:
            raise HTTPException(status_code=400, detail=f"版本 {version_id} 不存在")
        version_name = str(vrow[0])
        current_month = int(vrow[1] or 1)

        cur_data_accounts = await cdb.execute(
            """
            SELECT data_acct_code, data_acct_name, product_code, product_codes, value_type
            FROM data_account
            """
        )
        data_account_rows = await cur_data_accounts.fetchall()
        data_account_map = {
            str(r[0]): {
                "name": str(r[1]),
                "product_code": str(r[2]) if r[2] is not None else None,
                "value_type": str(r[4]),
            }
            for r in data_account_rows
        }

        cur_product = await cdb.execute(
            "SELECT product_code, product_name FROM product_type"
        )
        product_name_map = {
            str(r[0]): str(r[1]) for r in await cur_product.fetchall()
        }

        cur_period = await cdb.execute("SELECT period_id, year, month, quarter FROM period")
        period_map = {
            int(r[0]): {
                "year": str(r[1]),
                "month": str(r[2]),
                "quarter": str(r[3]),
            }
            for r in await cur_period.fetchall()
        }

        cur_report_accounts = await cdb.execute(
            "SELECT report_acct_code, report_acct_name, parent_code FROM report_account"
        )
        report_rows = await cur_report_accounts.fetchall()
        report_name_map = {str(r[0]): str(r[1]) for r in report_rows}
        report_parent_map = {
            str(r[0]): (str(r[2]) if r[2] is not None else None) for r in report_rows
        }

        cur_report_mappings = await cdb.execute(
            "SELECT report_acct_code, data_acct_code FROM report_data_mapping"
        )
        report_codes_by_data: dict[str, list[str]] = {}
        for report_code, data_code in await cur_report_mappings.fetchall():
            report_codes_by_data.setdefault(str(data_code), []).append(str(report_code))

        cur_dept_accounts = await cdb.execute(
            "SELECT dept_code, dept_name, parent_code FROM dept_account"
        )
        dept_rows = await cur_dept_accounts.fetchall()
        dept_name_map = {str(r[0]): str(r[1]) for r in dept_rows}
        dept_parent_map = {
            str(r[0]): (str(r[2]) if r[2] is not None else None) for r in dept_rows
        }

        cur_dept_mapping = await cdb.execute(
            "SELECT dept_code, product_code FROM dept_product_mapping"
        )
        dept_by_product = {str(r[1]): str(r[0]) for r in await cur_dept_mapping.fetchall()}

        cur_budget_data = await bdb.execute(
            """
            SELECT data_acct_code, product_code, period_id, budget_actual, value
            FROM budget_data
            WHERE version_id = ?
            """,
            (version_id,),
        )
        budget_rows = await cur_budget_data.fetchall()

        rows_to_insert: list[tuple[Any, ...]] = []
        update_time = _iso_now()
        for data_code_raw, product_code_raw, period_id_raw, budget_actual_raw, value_raw in budget_rows:
            data_code = str(data_code_raw)
            row_product_code = str(product_code_raw) if product_code_raw else ""
            period_id = int(period_id_raw)
            budget_actual = int(budget_actual_raw)
            value = float(value_raw or 0.0)
            acct = data_account_map.get(data_code)
            period = period_map.get(period_id)
            if not acct or not period:
                continue
            month_idx = _month_index(period["month"])
            expected_budget_actual = 1 if month_idx < current_month else 0
            if budget_actual != expected_budget_actual:
                continue

            product_code = row_product_code or acct.get("product_code") or ""
            product_name = product_name_map.get(product_code, "") if product_code else ""
            product_code_name = (
                f"{product_code} {product_name}".strip() if product_code else None
            )
            data_code_name = f"{data_code} {acct['name']}".strip()

            report_codes = sorted(set(report_codes_by_data.get(data_code, [])))
            if not report_codes:
                report_paths: list[list[str]] = [[]]
            else:
                report_paths = [
                    _build_report_path(code, report_name_map, report_parent_map)
                    for code in report_codes
                ]

            dept_code = dept_by_product.get(product_code) if product_code else None
            if dept_code:
                dept_path = _build_report_path(dept_code, dept_name_map, dept_parent_map)
            else:
                dept_path = []

            dept_levels = _build_fixed_levels(dept_path, 3)
            for report_path in report_paths:
                report_levels = _build_fixed_levels(report_path, 5)
                rows_to_insert.append(
                    (
                        report_levels[0],
                        report_levels[1],
                        report_levels[2],
                        report_levels[3],
                        report_levels[4],
                        dept_levels[0],
                        dept_levels[1],
                        dept_levels[2],
                        data_code_name,
                        product_code_name,
                        period["year"],
                        period["month"],
                        period["quarter"],
                        budget_actual,
                        version_id,
                        version_name,
                        value,
                        acct["value_type"],
                        update_time,
                    )
                )

        await bdb.execute("DELETE FROM budget_summary WHERE version_id = ?", (version_id,))
        if rows_to_insert:
            await bdb.executemany(
                """
                INSERT INTO budget_summary (
                  report_level1, report_level2, report_level3, report_level4, report_level5,
                  dept_level1, dept_level2, dept_level3, data_code_name, product_code_name,
                  year, month, quarter, budget_actual, version_id, version_name,
                  value, value_type, update_time
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                rows_to_insert,
            )
        await bdb.commit()
    return len(rows_to_insert)


app.include_router(
    build_budget_summary_compare_router(
        editable_context_provider=lambda: _editable_context(),
        get_version_current_month=lambda version_id, budget_path: _get_version_current_month(version_id, budget_path),
        rebuild_budget_summary_for_version=(
            lambda version_id, budget_path: _rebuild_budget_summary_for_version(version_id, budget_path)
        ),
        sync_compare_budget_summary=(
            lambda **kwargs: _sync_compare_budget_summary(**kwargs)
        ),
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
    )
)
app.include_router(
    build_compare_summary_export_router(
        compare_formula_export_jobs=_compare_formula_export_jobs,
        compare_formula_export_jobs_lock=_compare_formula_export_jobs_lock,
        run_compare_formula_export_job=lambda job_id, body: _run_compare_formula_export_job(job_id, body),
        export_compare_summary_full_pivot_callable=lambda: _export_compare_summary_full_pivot(),
    )
)
app.include_router(
    build_system_catalog_router(
        settings=settings,
        sync_databases_table_with_files=lambda: _sync_databases_table_with_files(),
        budget_schema=BUDGET_SCHEMA,
        get_year_period_months=lambda year: _get_year_period_months(year),
        iso_now=lambda: _iso_now(),
    )
)
app.include_router(
    build_system_admin_router(
        settings=settings,
        resolve_data_file_name=lambda data_file_id: _resolve_data_file_name(data_file_id),
        parse_year_from_budget_filename=lambda file_name: _parse_year_from_budget_filename(file_name),
        get_year_period_months=lambda year: _get_year_period_months(year),
        iso_now=lambda: _iso_now(),
        purge_disallowed_budget_data_for_version=(
            lambda bdb, version_id, current_month, period_month_map: _purge_disallowed_budget_data_for_version(
                bdb, version_id, current_month, period_month_map
            )
        ),
        sync_compare_budget_summary=lambda **kwargs: _sync_compare_budget_summary(**kwargs),
        validate_password_policy=lambda password: _validate_password_policy(password),
    )
)
app.include_router(
    build_system_edit_show_router(
        sync_compare_budget_summary=lambda **kwargs: _sync_compare_budget_summary(**kwargs),
    )
)
app.include_router(
    build_data_accounts_router(
        settings=settings,
        load_budget_data_ref_counts=lambda: _load_budget_data_ref_counts(),
        row_to_account=lambda row: _row_to_account(row),
        enrich_account_usage_flags=lambda account: _enrich_account_usage_flags(account),
        load_data_account_scope_map=lambda db: _load_data_account_scope_map(db),
        validate_formula_reference_scope=lambda **kwargs: _validate_formula_reference_scope(**kwargs),
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
        count_budget_data_refs=lambda code: _count_budget_data_refs(code),
        formulas_reference_code=lambda bf, af, code: formulas_reference_code(bf, af, code),
        get_account_row=lambda db, code: _get_account_row(db, code),
        normalize_cell=lambda value: _normalize_cell(value),
        color_row=lambda ws, row_idx, max_col, color: _color_row(ws, row_idx, max_col, color),
        normalize_formula=lambda formula: _normalize_formula(formula),
        latest_version=lambda: _latest_version(),
        try_latest_version_id=lambda: _try_latest_version_id(),
        recalculate_data_account_formula_all_products=(
            lambda **kwargs: _recalculate_data_account_formula_all_products(**kwargs)
        ),
        set_data_account_need_calc=lambda code, flag: _set_data_account_need_calc(code, flag),
        preview_insert_single_to_all_rows=(
            lambda code, old_product_code: preview_insert_single_to_all_rows(code, old_product_code)
        ),
        preview_delete_all_to_single_rows=(
            lambda code, keep_product_code: preview_delete_all_to_single_rows(code, keep_product_code)
        ),
        migrate_single_to_all_budget_data=(
            lambda code, old_product_code: migrate_single_to_all_budget_data(code, old_product_code)
        ),
        migrate_all_to_single_budget_data=(
            lambda code, keep_product_code: migrate_all_to_single_budget_data(code, keep_product_code)
        ),
    )
)
app.include_router(
    build_budget_subject_catalog_router(
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
    )
)
app.include_router(
    build_expense_actual_import_router(
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
    )
)
app.include_router(
    build_expense_forecast_router(
        default_year=settings.budget_year,
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
    )
)
app.include_router(
    build_report_catalog_router(
        normalize_cell=lambda value: _normalize_cell(value),
        color_row=lambda ws, row_idx, max_col, color: _color_row(ws, row_idx, max_col, color),
        validate_report_code_with_parent=(
            lambda code, level, parent_code: _validate_report_code_with_parent(code, level, parent_code)
        ),
        parse_bool_like=lambda value: _parse_bool_like(value),
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
    )
)
app.include_router(
    build_dept_catalog_router(
        normalize_cell=lambda value: _normalize_cell(value),
        color_row=lambda ws, row_idx, max_col, color: _color_row(ws, row_idx, max_col, color),
        validate_dept_code_with_parent=(
            lambda code, level, parent_code: _validate_dept_code_with_parent(code, level, parent_code)
        ),
        write_operation_log=lambda **kwargs: write_operation_log(**kwargs),
    )
)
app.include_router(
    build_budget_summary_export_router(
        editable_context_provider=lambda: _editable_context(),
        export_budget_summary_from_template=(
            lambda **kwargs: _export_budget_summary_from_template(**kwargs)
        ),
        export_budget_summary_formula_tree_workbook=(
            lambda body, version_id, budget_path, export_year: _export_budget_summary_formula_tree_workbook(
                body, version_id, budget_path, export_year
            )
        ),
    )
)
app.include_router(
    build_expense_budget_execution_router(
        editable_context_provider=lambda: _editable_context(),
    )
)


async def _sync_compare_budget_summary(
    trigger_source: str = "manual",
    operator_user_id: int | None = None,
) -> CompareSummarySyncResult:
    async with _compare_sync_lock:
        start_time = _iso_now()
        compare_path = compare_db_path()
        common_path = common_db_path()
        selected_versions = 0
        inserted_rows = 0
        message = "ok"
        rule_message = ""
        level_rules: list[str] = []
        status = "success"
        try:
            async with aiosqlite.connect(common_path) as cdb:
                await cdb.execute("PRAGMA foreign_keys = ON")
                cur = await cdb.execute(
                    """
                    SELECT e.edit_show_sign, e.data_file_id, e.version_id, d.data_file_name, d.year
                    FROM edit_show_version e
                    JOIN databases d ON d.id = e.data_file_id
                    WHERE e.edit_show_sign BETWEEN 1 AND 5
                    ORDER BY e.edit_show_sign
                    """
                )
                selected = await cur.fetchall()

            selected_versions = len(selected)
            rows_to_insert: list[tuple[Any, ...]] = []
            for show_level, data_file_id, source_version_id, data_file_name, source_year in selected:
                budget_path = settings.data_dir / str(data_file_name)
                if not budget_path.exists():
                    continue
                async with aiosqlite.connect(budget_path) as bdb:
                    await bdb.execute("PRAGMA foreign_keys = ON")
                    cur_ver = await bdb.execute(
                        "SELECT version_name, current_month FROM version WHERE version_id = ?",
                        (int(source_version_id),),
                    )
                    ver_row = await cur_ver.fetchone()
                    ver_name = str(ver_row[0]) if ver_row and ver_row[0] is not None else None
                    current_month = int(ver_row[1] or 1) if ver_row else 1
                    level_rules.append(
                        f"L{int(show_level)}: {data_file_name} / V{int(source_version_id)}"
                        f"{(' ' + ver_name) if ver_name else ''} / current_month={current_month}"
                        f" / month<{current_month}取实际, month>={current_month}取预算"
                    )
                    # 若展示版本的 budget_summary 尚未生成（或被清空），先补一次重建，避免 compare 丢年度。
                    cur_summary_cnt = await bdb.execute(
                        "SELECT COUNT(*) FROM budget_summary WHERE version_id = ?",
                        (int(source_version_id),),
                    )
                    summary_cnt_row = await cur_summary_cnt.fetchone()
                    if int(summary_cnt_row[0] or 0) == 0:
                        await _rebuild_budget_summary_for_version(
                            int(source_version_id),
                            budget_path,
                        )
                        cur_summary_cnt = await bdb.execute(
                            "SELECT COUNT(*) FROM budget_summary WHERE version_id = ?",
                            (int(source_version_id),),
                        )
                        summary_cnt_row = await cur_summary_cnt.fetchone()
                    if int(summary_cnt_row[0] or 0) == 0:
                        level_rules.append(
                            f"L{int(show_level)}: {data_file_name} / V{int(source_version_id)} 汇总明细为空（budget_data 无记录）"
                        )
                    cur = await bdb.execute(
                        """
                        SELECT report_level1, report_level2, report_level3, report_level4, report_level5,
                               dept_level1, dept_level2, dept_level3, data_code_name, product_code_name,
                               year, month, quarter, budget_actual, version_name, value, value_type
                        FROM budget_summary
                        WHERE version_id = ?
                        ORDER BY report_level1, report_level2, report_level3, data_code_name, month, budget_actual
                        """,
                        (int(source_version_id),),
                    )
                    for row in await cur.fetchall():
                        rows_to_insert.append(
                            (
                                int(show_level),
                                int(data_file_id),
                                int(source_year),
                                int(source_version_id),
                                ver_name or row[14],  # source_version_name
                                row[0],
                                row[1],
                                row[2],
                                row[3],
                                row[4],
                                row[5],
                                row[6],
                                row[7],
                                row[8],
                                row[9],
                                row[10],
                                row[11],
                                row[12],
                                int(row[13] or 0),
                                float(row[15] or 0.0),
                                row[16],
                                start_time,
                            )
                        )

            async with aiosqlite.connect(compare_path) as cdb:
                await cdb.execute("PRAGMA foreign_keys = ON")
                await cdb.execute("DELETE FROM compare_budget_summary")
                if rows_to_insert:
                    await cdb.executemany(
                        """
                        INSERT INTO compare_budget_summary (
                          show_level, data_file_id, source_year, source_version_id, source_version_name,
                          report_level1, report_level2, report_level3, report_level4, report_level5,
                          dept_level1, dept_level2, dept_level3, data_code_name, product_code_name,
                          year, month, quarter, budget_actual, value, value_type, sync_time
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        rows_to_insert,
                    )
                inserted_rows = len(rows_to_insert)
                if level_rules:
                    rule_message = "；".join(level_rules)
                end_time = _iso_now()
                await cdb.execute(
                    """
                    INSERT INTO compare_sync_job_log (
                      start_time, end_time, trigger_source, status, message, operator_user_id
                    ) VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (start_time, end_time, trigger_source, status, message, operator_user_id),
                )
                await cdb.commit()
            await _set_compare_refresh_time_b(end_time)
        except Exception as exc:
            status = "failed"
            message = str(exc)
            end_time = _iso_now()
            async with aiosqlite.connect(compare_path) as cdb:
                await cdb.execute("PRAGMA foreign_keys = ON")
                await cdb.execute(
                    """
                    INSERT INTO compare_sync_job_log (
                      start_time, end_time, trigger_source, status, message, operator_user_id
                    ) VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (start_time, end_time, trigger_source, status, message, operator_user_id),
                )
                await cdb.commit()
            raise HTTPException(status_code=500, detail=f"同步 compare_summary 失败: {exc}")

        return CompareSummarySyncResult(
            inserted_rows=inserted_rows,
            selected_versions=selected_versions,
            trigger_source=trigger_source,
            message=message,
            rule_message=rule_message,
            level_rules=level_rules,
        )


def _parse_year_from_budget_filename(name: str) -> int | None:
    m = re.match(r"budget_(\d{4})\.db$", name)
    if not m:
        return None
    return int(m.group(1))


def _fmt_file_ctime(path: Path) -> str:
    try:
        return datetime.fromtimestamp(path.stat().st_ctime, timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    except Exception:
        return _iso_now()


def _validate_password_policy(password: str) -> None:
    text = password or ""
    if len(text) < 8:
        raise HTTPException(status_code=400, detail="密码至少 8 位")
    has_alpha = any(ch.isalpha() for ch in text)
    if not has_alpha:
        raise HTTPException(status_code=400, detail="密码至少包含 1 个字母，且区分大小写")


async def _resolve_data_file_name(data_file_id: int) -> str:
    async with aiosqlite.connect(common_db_path()) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            "SELECT data_file_name FROM databases WHERE id = ?",
            (data_file_id,),
        )
        row = await cur.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail=f"data_file_id={data_file_id} 不存在")
        return str(row[0])


async def _sync_databases_table_with_files() -> list[SystemDatabaseRow]:
    data_dir = settings.data_dir
    data_dir.mkdir(parents=True, exist_ok=True)
    files = sorted([p for p in data_dir.glob("budget_*.db") if p.is_file()])
    seen_names = {p.name for p in files}

    async with aiosqlite.connect(common_db_path()) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute("SELECT id, data_file_name FROM databases")
        db_rows = await cur.fetchall()
        db_name_to_id = {str(r[1]): int(r[0]) for r in db_rows}

        # 删除文件系统不存在的记录
        for file_name, data_file_id in db_name_to_id.items():
            if file_name not in seen_names:
                await db.execute("DELETE FROM edit_show_version WHERE data_file_id = ?", (data_file_id,))
                await db.execute("DELETE FROM databases WHERE id = ?", (data_file_id,))

        # 新增或刷新存在的文件
        for fp in files:
            file_name = fp.name
            year = _parse_year_from_budget_filename(file_name) or settings.budget_year
            create_time = _fmt_file_ctime(fp)
            try:
                async with aiosqlite.connect(fp) as bdb:
                    cur_settings = await bdb.execute(
                        "SELECT setting_key, setting_value FROM settings WHERE setting_key IN ('year', 'create_time')"
                    )
                    settings_rows = await cur_settings.fetchall()
                    settings_map = {str(r[0]): str(r[1]) for r in settings_rows}
                    if settings_map.get("year", "").isdigit():
                        year = int(settings_map["year"])
                    if settings_map.get("create_time"):
                        create_time = settings_map["create_time"]
            except Exception:
                pass

            await db.execute(
                """
                INSERT INTO databases(data_file_name, year, create_time)
                VALUES (?, ?, ?)
                ON CONFLICT(data_file_name) DO UPDATE SET
                  year = excluded.year,
                  create_time = excluded.create_time
                """,
                (file_name, int(year), create_time),
            )

        await db.commit()
        cur = await db.execute(
            "SELECT id, data_file_name, year, create_time FROM databases ORDER BY year DESC, data_file_name"
        )
        rows = await cur.fetchall()
    return [
        SystemDatabaseRow(
            id=int(r[0]),
            data_file_name=str(r[1]),
            year=int(r[2]),
            create_time=str(r[3]),
            file_path=str(settings.data_dir / str(r[1])),
        )
        for r in rows
    ]


async def _get_year_period_months(year: int) -> dict[int, int]:
    out: dict[int, int] = {}
    year_label = f"Y{year}"
    async with aiosqlite.connect(common_db_path()) as cdb:
        await cdb.execute("PRAGMA foreign_keys = ON")
        cur = await cdb.execute(
            """
            SELECT period_id, month
            FROM period
            WHERE year = ?
            ORDER BY period_id
            """,
            (year_label,),
        )
        for period_id, month_label in await cur.fetchall():
            m = _month_index(str(month_label or ""))
            if 1 <= m <= 12:
                out[int(period_id)] = m
    return out


async def _fetch_budget_summary_all_rows(
    budget_path: Path | None = None,
) -> list[tuple[Any, ...]]:
    path = budget_path if budget_path is not None else budget_db_path(settings.budget_year)
    async with aiosqlite.connect(path) as db:
        cur = await db.execute(
            """
            SELECT report_level1, report_level2, report_level3, report_level4, report_level5,
                   dept_level1, dept_level2, dept_level3, data_code_name, product_code_name,
                   year, month, quarter, budget_actual, version_id, version_name,
                   value, value_type, update_time
            FROM budget_summary
            ORDER BY version_id, report_level1, report_level2, report_level3, data_code_name, month, budget_actual
            """
        )
        return await cur.fetchall()


def _extract_data_acct_code_from_name(data_code_name: str) -> str | None:
    m = re.match(r"\s*([A-Z]\d{4})\b", (data_code_name or "").strip())
    return m.group(1) if m else None


def _extract_product_code_from_summary_name(product_code_name: str | None) -> str | None:
    if not product_code_name:
        return None
    s = str(product_code_name).strip()
    if not s:
        return None
    return s.split()[0]


def _formula_ref_display_value(raw_value: float, value_type: str | None) -> str:
    if (value_type or "") == "百分比":
        display = raw_value if abs(raw_value) > 1 else raw_value * 100.0
        return f"{display:,.1f}%"
    return f"{raw_value:,.1f}"


async def _export_budget_summary_from_template(
    body: BudgetSummaryExportPivotRequest,
    output_filename: str,
):
    return await _get_budget_summary_export_service().export_budget_summary_from_template(
        body=body,
        output_filename=output_filename,
    )


def _month_idx_from_label(label: str | None) -> int | None:
    text = (label or "").strip().upper()
    m = re.search(r"(\d{1,2})", text)
    if not m:
        return None
    month = int(m.group(1))
    if month < 1 or month > 12:
        return None
    return month - 1


async def _export_budget_summary_formula_tree_workbook(
    _body: BudgetSummaryExportPivotRequest,
    version_id: int,
    budget_path: Path,
    budget_year: int,
) -> StreamingResponse:
    return await _get_budget_summary_export_service().export_budget_summary_formula_tree_workbook(
        _body=_body,
        version_id=version_id,
        budget_path=budget_path,
        budget_year=budget_year,
    )


_budget_summary_export_service: BudgetSummaryExportService | None = None


def _get_budget_summary_export_service() -> BudgetSummaryExportService:
    global _budget_summary_export_service
    if _budget_summary_export_service is None:
        _budget_summary_export_service = BudgetSummaryExportService(
            editable_context_provider=lambda: _editable_context(),
            fetch_budget_summary_all_rows=lambda budget_path: _fetch_budget_summary_all_rows(budget_path),
            normalize_formula=lambda formula: _normalize_formula(formula),
            normalize_summary_value=lambda field_id, raw: normalize_summary_value(field_id, raw),
            extract_data_acct_code_from_name=lambda data_code_name: _extract_data_acct_code_from_name(data_code_name),
            extract_product_code_from_summary_name=lambda product_code_name: _extract_product_code_from_summary_name(
                product_code_name
            ),
            formula_ref_display_value=lambda value, value_type: _formula_ref_display_value(value, value_type),
            budget_summary_field_meta=lambda: budget_summary_field_meta(),
            write_template_pivot_data_area=lambda **kwargs: write_template_pivot_data_area(**kwargs),
            build_export_year_datetime_text=lambda export_years: build_export_year_datetime_text(export_years),
            month_idx_from_label=lambda label: _month_idx_from_label(label),
            prepare_formula_expression=lambda formula: _prepare_formula_expression(formula),
            autosize_worksheet_columns=lambda ws: autosize_worksheet_columns(ws),
        )
    return _budget_summary_export_service


_compare_export_service: CompareExportService | None = None


def _get_compare_export_service() -> CompareExportService:
    global _compare_export_service
    if _compare_export_service is None:
        _compare_export_service = CompareExportService(
            data_dir=settings.data_dir,
            compare_formula_export_jobs=_compare_formula_export_jobs,
            compare_formula_export_jobs_lock=_compare_formula_export_jobs_lock,
            sync_compare_budget_summary=lambda **kwargs: _sync_compare_budget_summary(**kwargs),
            extract_data_acct_code_from_name=lambda data_code_name: _extract_data_acct_code_from_name(data_code_name),
            extract_product_code_from_summary_name=lambda product_code_name: _extract_product_code_from_summary_name(
                product_code_name
            ),
            month_idx_from_label=lambda label: _month_idx_from_label(label),
            normalize_formula=lambda formula: _normalize_formula(formula),
            prepare_formula_expression=lambda formula: _prepare_formula_expression(formula),
            autosize_worksheet_columns=lambda ws: autosize_worksheet_columns(ws),
            normalize_summary_value=lambda field_id, raw: normalize_summary_value(field_id, raw),
            budget_summary_field_meta=lambda: budget_summary_field_meta(),
            write_template_pivot_data_area=lambda **kwargs: write_template_pivot_data_area(**kwargs),
            build_export_versions_info_text=lambda version_rows: build_export_versions_info_text(version_rows),
        )
    return _compare_export_service


async def _run_compare_formula_export_job(job_id: str, body: BudgetSummaryExportPivotRequest) -> None:
    await _get_compare_export_service().run_compare_formula_export_job(job_id, body)


async def _export_compare_summary_full_pivot():
    return await _get_compare_export_service().export_compare_summary_full_pivot()


def _row_to_account(r: tuple[Any, ...]) -> DataAccountRow:
    # 新字段顺序：data_acct_code, data_acct_name, metric_group_code, metric_group_name,
    #              product_code, product_codes, budget_formula, actual_formula,
    #              need_calc, value_type, remark, product_name
    product_codes_val = r[5] if len(r) > 5 else None
    return DataAccountRow(
        data_acct_code=r[0],
        data_acct_name=r[1],
        metric_group_code=r[2] if len(r) > 2 else None,
        metric_group_name=r[3] if len(r) > 3 else None,
        product_code=r[4] if len(r) > 4 else None,
        product_codes=product_codes_val,
        budget_formula=r[6] if len(r) > 6 else None,
        actual_formula=r[7] if len(r) > 7 else None,
        need_calc=int(r[8] or 0) if len(r) > 8 else 0,
        value_type=r[9] if len(r) > 9 else "金额",
        remark=r[10] if len(r) > 10 else None,
        product_name=r[11] if len(r) > 11 else None,
        has_budget_data_records=False,
    )


async def _load_budget_data_ref_counts() -> dict[str, int]:
    counts: dict[str, int] = {}
    for path in list_budget_database_files():
        if not path.exists():
            continue
        async with aiosqlite.connect(path) as db:
            await db.execute("PRAGMA foreign_keys = ON")
            cur = await db.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='budget_data'"
            )
            if not await cur.fetchone():
                continue
            cur = await db.execute(
                """
                SELECT data_acct_code, COUNT(*)
                FROM budget_data
                WHERE data_acct_code IS NOT NULL
                GROUP BY data_acct_code
                """
            )
            for r in await cur.fetchall():
                if not r[0]:
                    continue
                c = str(r[0])
                counts[c] = counts.get(c, 0) + int(r[1] or 0)
    return counts


async def _count_budget_data_refs(code: str) -> int:
    total = 0
    for path in list_budget_database_files():
        if not path.exists():
            continue
        async with aiosqlite.connect(path) as db:
            await db.execute("PRAGMA foreign_keys = ON")
            cur = await db.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='budget_data'"
            )
            if not await cur.fetchone():
                continue
            cur = await db.execute(
                "SELECT COUNT(*) FROM budget_data WHERE data_acct_code = ?",
                (code,),
            )
            total += int((await cur.fetchone())[0] or 0)
    return total


async def _count_report_mapping_refs(code: str) -> int:
    path = common_db_path()
    async with aiosqlite.connect(path) as db:
        await db.execute("PRAGMA foreign_keys = ON")
        cur = await db.execute(
            "SELECT COUNT(*) FROM report_data_mapping WHERE data_acct_code = ?",
            (code,),
        )
        return int((await cur.fetchone())[0] or 0)


async def _enrich_account_usage_flags(account: DataAccountRow) -> DataAccountRow:
    budget_ref_count = await _count_budget_data_refs(account.data_acct_code)
    mapping_ref_count = await _count_report_mapping_refs(account.data_acct_code)
    account.budget_data_ref_count = budget_ref_count
    account.report_mapping_ref_count = mapping_ref_count
    account.has_budget_data_records = budget_ref_count > 0
    return account


async def _get_account_row(
    db: aiosqlite.Connection, code: str
) -> dict[str, Any] | None:
    cur = await db.execute(
        """
        SELECT data_acct_code, data_acct_name, product_code, product_codes,
               budget_formula, actual_formula, need_calc, value_type, remark
        FROM data_account WHERE data_acct_code = ?
        """,
        (code,),
    )
    r = await cur.fetchone()
    if not r:
        return None
    return {
        "data_acct_code": r[0],
        "data_acct_name": r[1],
        "product_code": r[2],
        "product_codes": r[3],
        "budget_formula": r[4],
        "actual_formula": r[5],
        "need_calc": int(r[6] or 0),
        "value_type": r[7],
        "remark": r[8],
    }


def _normalize_cell(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _color_row(ws: Any, row_idx: int, max_col: int, color: str) -> None:
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_idx, column=col)
        base_font = cell.font
        cell.font = Font(
            name=base_font.name,
            size=base_font.size,
            bold=base_font.bold,
            italic=base_font.italic,
            underline=base_font.underline,
            strike=base_font.strike,
            color=color,
        )


def _validate_dept_code_with_parent(code: str, level: int, parent_code: str | None) -> str | None:
    if level == 1:
        if not re.match(r"^Y\d{1,2}$", code):
            return "1级部门科目代码格式错误（示例：Y1 或 Y01）"
        return None
    if not parent_code:
        return f"缺少上级部门科目，无法校验第{level}级部门科目代码"
    if not code.startswith(parent_code):
        return f"第{level}级部门科目代码必须继承上级代码前缀"
    suffix = code[len(parent_code):]
    if not re.match(r"^\d{1,2}$", suffix):
        return f"第{level}级部门科目代码应为“上级代码 + 1-2位数字”"
    return None


def _validate_report_code_with_parent(code: str, level: int, parent_code: str | None) -> str | None:
    if level == 1:
        if not re.match(r"^[A-Z]\d{2}$", code):
            return "1级报告科目代码格式错误（示例：A01）"
        return None
    if not parent_code:
        return f"缺少上级报告科目，无法校验第{level}级报告科目代码"
    if not code.startswith(parent_code):
        return f"第{level}级报告科目代码必须继承上级代码前缀"
    suffix = code[len(parent_code):]
    if not re.match(r"^\d{2}$", suffix):
        return f"第{level}级报告科目代码应为“上级代码 + 2位数字”"
    return None


def _parse_bool_like(v: str) -> int | None:
    s = (v or "").strip().lower()
    if not s:
        return None
    if s in {"1", "true", "yes", "y", "是", "对"}:
        return 1
    if s in {"0", "false", "no", "n", "否", "错"}:
        return 0
    return None


